from pathlib import Path

from openpyxl import Workbook
import pytest

from src.mass_calculator import calculate_total_mass


def _grid_star(node_id: int, x: float, y: float, z: float) -> str:
    first = f"{'GRID*':<8}{node_id:<16}{'':<16}{x:<16.8f}{y:<16.8f}"
    second = f"{'*':<8}{z:<16.8f}"
    return f"{first}\n{second}"


def test_mass_summary_includes_budget_comparison_and_region_classification(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mass Budget"
    values_grams = {
        7: 2.0,
        8: 0.0,
        9: 0.0,
        10: 0.0,
        11: 1.0,
        12: 3.0,
    }
    labels = {
        7: "H-PSU module",
        8: "MOD module",
        9: "FPGA module",
        10: "PROC module",
        11: "Backplane modules",
        12: "Total mass",
    }
    for row, label in labels.items():
        sheet.cell(row=row, column=2, value=label)
        sheet.cell(row=row, column=3, value=values_grams[row])
        sheet.cell(row=row, column=5, value=values_grams[row])
    budget_path = tmp_path / "budget.xlsx"
    workbook.save(budget_path)

    bdf_text = "\n".join(
        [
            "$ Elements and Element Properties for region : 09_A_t-boards-PSU",
            "PSHELL 25 2 0.001 2 2",
            "CQUAD4 1 25 1 2 3 4 0. 0.",
            "$ Elements and Element Properties for region : 07_t-bandejas",
            "PSHELL 24 1 0.001 1 1",
            "CQUAD4 2 24 5 6 7 8 0. 0.",
            "MAT1,1,1.0,,0.3,1.0",
            "MAT1,2,1.0,,0.3,2.0",
            _grid_star(1, 0.0, 0.0, 0.00815),
            _grid_star(2, 1.0, 0.0, 0.00815),
            _grid_star(3, 1.0, 1.0, 0.00815),
            _grid_star(4, 0.0, 1.0, 0.00815),
            _grid_star(5, 0.0, 0.0, 0.0),
            _grid_star(6, 1.0, 0.0, 0.0),
            _grid_star(7, 1.0, 1.0, 0.0),
            _grid_star(8, 0.0, 1.0, 0.0),
            "ENDDATA",
        ]
    )
    bdf_path = tmp_path / "sample.bdf"
    bdf_path.write_text(bdf_text, encoding="utf-8")

    summary = calculate_total_mass(
        bdf_path=bdf_path,
        target_mass=0.003,
        budget_workbook=budget_path,
        budget_sheet="Mass Budget",
        budget_basis="basic",
    )

    assert summary.total_mass == pytest.approx(0.003)
    assert summary.relative_error == pytest.approx(0.0)
    assert summary.budget_comparison is not None
    assert summary.budget_comparison.basis == "basic"

    rows = {row.category: row for row in summary.budget_comparison.rows}
    assert rows["H-PSU module"].model_mass == pytest.approx(0.002)
    assert rows["Backplane modules"].model_mass == pytest.approx(0.001)
    assert rows["Total mass"].delta_mass == pytest.approx(0.0)

    classifications = {
        (row.region_name, row.assigned_to): row.mass for row in summary.budget_comparison.classification_rows
    }
    assert classifications[("09_A_t-boards-PSU", "H-PSU module")] == pytest.approx(0.002)
    assert classifications[("07_t-bandejas", "Backplane modules")] == pytest.approx(0.001)
